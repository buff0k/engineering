# Copyright (c) 2026, BuFf0k and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
from frappe.utils import now_datetime, nowdate

from engineering.controllers.fleet_compliance import (
	apply_draft_allocation_penalty,
	bulk_drivers,
	collect_drivers,
	compute_all,
	effective_asset_owner,
	get_effective_allocations,
	get_expiring_threshold_days,
	render_addendum_html,
	render_driver_licence_html,
	render_drivers_overview_html,
	render_service_history_html,
	render_vehicle_licence_html,
)
from engineering.engineering.doctype.fleet_management_settings.fleet_management_settings import (
	get_location_custodian,
)


@frappe.whitelist()
@frappe.validate_and_sanitize_search_inputs
def public_road_asset_query(doctype, txt, searchfield, start, page_len, filters):
	"""Link query for the Asset field on both Vehicle Allocation (asset) and
	Vehicle Licence (fleet_number): only submitted Assets whose Category is
	listed on Fleet Management Settings -> Public Road Asset Categories.
	Submitted-only because a Draft Asset isn't a real in-service vehicle
	yet and a Cancelled one no longer is (same rule the Vehicle Licence
	Expiration page's counts and the road asset register export use)."""
	from engineering.engineering.doctype.fleet_management_settings.fleet_management_settings import (
		get_public_road_asset_categories,
	)

	categories = get_public_road_asset_categories()

	if not categories:
		return []

	return frappe.db.sql(
		"""
		select name, asset_name
		from `tabAsset`
		where asset_category in %(categories)s
			and docstatus = 1
			and (name like %(txt)s or asset_name like %(txt)s)
		order by name
		limit %(start)s, %(page_len)s
		""",
		{
			"categories": categories,
			"txt": f"%{txt}%",
			"start": start,
			"page_len": page_len,
		},
	)


@frappe.whitelist()
@frappe.validate_and_sanitize_search_inputs
def allocated_asset_query(doctype, txt, searchfield, start, page_len, filters):
	"""Link query for the Vehicle Allocation list view's Asset quick filter —
	unlike public_road_asset_query (used on the form, where you're picking an
	Asset to *create* a new allocation for), this only offers Assets that are
	both a public-road category AND already have at least one Vehicle
	Allocation. Anything else would just filter the list down to zero rows."""
	from engineering.engineering.doctype.fleet_management_settings.fleet_management_settings import (
		get_public_road_asset_categories,
	)

	categories = get_public_road_asset_categories()

	if not categories:
		return []

	return frappe.db.sql(
		"""
		select distinct a.name, a.asset_name
		from `tabAsset` a
		where a.asset_category in %(categories)s
			and (a.name like %(txt)s or a.asset_name like %(txt)s)
			and exists (
				select 1 from `tabVehicle Allocation` va where va.asset = a.name
			)
		order by a.name
		limit %(start)s, %(page_len)s
		""",
		{
			"categories": categories,
			"txt": f"%{txt}%",
			"start": start,
			"page_len": page_len,
		},
	)


@frappe.whitelist()
@frappe.validate_and_sanitize_search_inputs
def allocation_driver_query(doctype, txt, searchfield, start, page_len, filters):
	"""Link query for the Vehicle Allocation list view's Driver quick filter —
	only offers Employees who actually appear as a driver on at least one
	Vehicle Allocation, same reasoning as allocated_asset_query above."""
	return frappe.db.sql(
		"""
		select distinct e.name, e.employee_name
		from `tabEmployee` e
		where (e.name like %(txt)s or e.employee_name like %(txt)s)
			and exists (
				select 1 from `tabVehicle Allocation Driver` vad where vad.driver = e.name
			)
		order by e.employee_name
		limit %(start)s, %(page_len)s
		""",
		{"txt": f"%{txt}%", "start": start, "page_len": page_len},
	)


@frappe.whitelist()
def get_overall_statuses(names):
	"""overall_status (and the fields it depends on) is a virtual field, so
	it never comes through the List View's bulk query. The list view JS
	calls this once per page of rows and patches the indicator in
	afterwards, rather than storing/caching the value anywhere."""
	if isinstance(names, str):
		names = frappe.parse_json(names)

	if not names:
		return {}

	threshold_days = get_expiring_threshold_days()

	rows = frappe.get_all(
		"Vehicle Allocation",
		filters={"name": ["in", names]},
		fields=["name", "asset", "required_licence_type"],
	)

	drivers_by_parent = bulk_drivers([row.name for row in rows])

	return {
		row.name: compute_all(
			row.asset,
			[d.driver for d in drivers_by_parent.get(row.name, [])],
			row.required_licence_type,
			threshold_days,
		)["overall_status"]
		for row in rows
	}


@frappe.whitelist()
def preview_compliance(asset=None, required_licence_type=None, drivers=None):
	"""Lets the form recompute every compliance HTML block (plus Overall
	Status) the instant Asset / Drivers / Required Licence Type changes,
	without needing a save + reload first. Takes raw values rather than a
	saved doc name on purpose — the whole point is to preview a combination
	that may not be saved yet (e.g. the user is mid-edit on an existing
	allocation)."""
	if not frappe.has_permission("Vehicle Allocation", "read"):
		frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

	if isinstance(drivers, str):
		drivers = frappe.parse_json(drivers)

	drivers = collect_drivers(drivers)

	return {
		"vehicle_licence_compliance_html": render_vehicle_licence_html(asset),
		"driver_licence_compliance_html": render_driver_licence_html(drivers, required_licence_type),
		"company_vehicle_undertaking_html": render_addendum_html(drivers),
		"drivers_overview_html": render_drivers_overview_html(drivers, required_licence_type),
		"service_history_html": render_service_history_html(asset),
		"overall_status": compute_all(asset, drivers, required_licence_type)["overall_status"],
	}


def _find_active_conflicts(asset, drivers, exclude_name=None):
	"""Two kinds of overlap worth flagging before this allocation is
	submitted — both resolved automatically on submit, mirroring each
	other:

	1. "asset" — this Asset is already the subject of another still-open
	   ("Current", submitted) allocation. close_previous_open_allocation
	   closes the other one.
	2. "employee" — a Driver listed here already drives a DIFFERENT Asset
	   under another still-open allocation — a change of vehicle for that
	   driver. close_or_update_previous_driver_allocations closes the
	   other allocation too, UNLESS it's a shared vehicle (more than one
	   Driver) — closing the whole thing just because one of several
	   drivers moved on would strand the rest, so there this driver is
	   only removed from it, not the whole allocation. sole_driver on the
	   returned dict says which case applies, so callers can word the
	   heads-up accordingly.

	Used both by the whitelisted client-facing wrapper and by validate()
	(server-side, so the warning surfaces on every save regardless of what
	the browser's JS did or didn't run)."""
	exclude_name = exclude_name or ""
	drivers = collect_drivers(drivers or [])
	conflicts = []

	if asset:
		for row in frappe.get_all(
			"Vehicle Allocation",
			filters={"asset": asset, "status": "Current", "docstatus": 1, "name": ["!=", exclude_name]},
			fields=["name"],
		):
			conflicts.append({"type": "asset", "allocation": row.name})

	if drivers:
		driver_rows = frappe.get_all(
			"Vehicle Allocation Driver",
			filters={"driver": ["in", drivers], "parenttype": "Vehicle Allocation", "parentfield": "drivers"},
			fields=["driver", "driver_name", "parent"],
		)
		candidate_parents = {r.parent for r in driver_rows} - {exclude_name}

		if candidate_parents:
			open_parents = {
				a.name: a
				for a in frappe.get_all(
					"Vehicle Allocation",
					filters={"name": ["in", list(candidate_parents)], "status": "Current", "docstatus": 1},
					fields=["name", "asset"],
				)
			}

			all_drivers_by_parent = {}

			for r in frappe.get_all(
				"Vehicle Allocation Driver",
				filters={"parent": ["in", list(candidate_parents)], "parentfield": "drivers"},
				fields=["parent", "driver"],
			):
				all_drivers_by_parent.setdefault(r.parent, set()).add(r.driver)

			for row in driver_rows:
				parent = open_parents.get(row.parent)

				if not parent or parent.asset == asset:
					# Not open, or it's this same Asset — already covered by
					# the "asset" conflict above, not a second finding.
					continue

				other_drivers = all_drivers_by_parent.get(row.parent, set()) - {row.driver}

				conflicts.append(
					{
						"type": "employee",
						"driver": row.driver,
						"driver_name": row.driver_name,
						"allocation": row.parent,
						"other_asset": parent.asset,
						"sole_driver": not other_drivers,
					}
				)

	return conflicts


@frappe.whitelist()
def check_active_conflicts(asset=None, drivers=None, exclude_name=None):
	"""Client-facing wrapper around _find_active_conflicts — called from the
	form on load and whenever Asset/Drivers change, so the warning shows up
	before the user even attempts to save."""
	if not frappe.has_permission("Vehicle Allocation", "read"):
		frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

	if isinstance(drivers, str):
		drivers = frappe.parse_json(drivers)

	return _find_active_conflicts(asset, drivers, exclude_name)


@frappe.whitelist()
def export_road_asset_register_xlsx():
	"""XLSX export of every submitted public-road Asset (registered or not),
	grouped by owner — same general shape as the manually-maintained LDV
	spreadsheet this replaces, but Location/Driver/compliance are pulled
	live from the actual Vehicle Allocation records instead of free-typed
	site codes and driver names. Draft/Cancelled Assets are excluded — same
	rule used everywhere else in fleet compliance.

	"Owner" is not simply the Asset's Company field — that's just the books
	the Asset is accounted under. A hired/leased LDV has Asset Owner =
	Supplier, so its actual owner is its Supplier, not whichever Company it
	happens to be booked against. See effective_asset_owner().

	A Draft Vehicle Allocation still populates Location/Drivers/etc (it
	shouldn't make the Asset look unregistered just because nobody has
	submitted it yet) — but its Overall Status always reads Non-Compliant
	regardless of how the underlying details look, since the allocation
	itself isn't legally in effect yet. See get_effective_allocations."""
	if not frappe.has_permission("Vehicle Allocation", "read"):
		frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

	from engineering.engineering.doctype.fleet_management_settings.fleet_management_settings import (
		get_reportable_asset_names,
	)

	asset_names = get_reportable_asset_names()

	if asset_names is None:
		frappe.throw(frappe._("No Public Road Asset Categories are configured on Fleet Management Settings."))

	if not asset_names:
		frappe.throw(frappe._("No Assets found within the configured Reporting Scope."))

	rows = frappe.db.sql(
		"""
		select
			a.name as asset,
			a.asset_name as asset_name,
			a.item_name as model,
			a.company as company,
			a.asset_owner as asset_owner,
			a.supplier as supplier,
			a.customer as customer
		from `tabAsset` a
		where a.name in %(asset_names)s
		order by a.name asc
		""",
		{"asset_names": list(asset_names)},
		as_dict=True,
	)

	for row in rows:
		row["owner"] = effective_asset_owner(row.asset_owner, row.company, row.supplier, row.customer)

	allocation_by_asset = get_effective_allocations(
		[r.asset for r in rows], fields=["location", "required_licence_type", "comments"]
	)

	for row in rows:
		effective = allocation_by_asset.get(row.asset)
		row["allocation"] = effective.name if effective else None
		row["location"] = effective.location if effective else None
		row["required_licence_type"] = effective.required_licence_type if effective else None
		row["comments"] = effective.comments if effective else None

	# Bulk-fetch the latest Vehicle Licence per Asset for the registration/
	# chassis/engine/VIN detail columns — one query, not N. Unlike compliance
	# *status* (strictly submitted-only, by design), these are just facts
	# about the physical vehicle, so a Draft record still counts here —
	# submitted wins over draft on a tie, otherwise the newest issue_date
	# wins regardless of docstatus.
	licence_by_asset = {}
	for lic in frappe.get_all(
		"Vehicle Licence",
		filters={"fleet_number": ["in", [r.asset for r in rows]], "docstatus": ["<", 2]},
		fields=["fleet_number", "registration_number", "chassis", "engine_number", "vin_number", "docstatus"],
		order_by="issue_date desc, docstatus desc",
	):
		licence_by_asset.setdefault(lic.fleet_number, lic)  # first hit per asset = latest (ordered desc)

	threshold_days = get_expiring_threshold_days()
	drivers_by_allocation = bulk_drivers([r.allocation for r in rows if r.allocation])

	for row in rows:
		driver_rows = drivers_by_allocation.get(row.allocation, [])
		row.update(
			compute_all(
				row.asset,
				[d.driver for d in driver_rows],
				row.required_licence_type,
				threshold_days,
			)
		)

		effective = allocation_by_asset.get(row.asset)
		if effective:
			row["overall_status"] = apply_draft_allocation_penalty(row["overall_status"], effective.docstatus)

		lic = licence_by_asset.get(row.asset)
		row["registration_number"] = lic.registration_number if lic else None
		row["chassis"] = lic.chassis if lic else None
		row["engine_number"] = lic.engine_number if lic else None
		row["vin_number"] = lic.vin_number if lic else None
		row["driver_display"] = ", ".join(
			f"{d.driver} - {d.driver_name}" if d.driver_name else d.driver for d in driver_rows
		)

	return {
		"filename": f"road_asset_register_{nowdate()}.xlsx",
		"content": _rows_to_xlsx_base64(rows),
		"type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
	}


def _sanitize_table_name(label, index):
	"""Excel table (ListObject) names must be unique in the workbook, start
	with a letter, and contain only letters/digits/underscore."""
	cleaned = "".join(ch if ch.isalnum() else "_" for ch in (label or ""))

	if not cleaned or not cleaned[0].isalpha():
		cleaned = f"Owner_{cleaned}"

	return f"{cleaned}_{index}"[:255]


def _rows_to_xlsx_base64(rows):
	import base64
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill
	from openpyxl.worksheet.table import Table, TableStyleInfo

	wb = Workbook()
	ws = wb.active
	ws.title = "Road Asset Register"

	columns = [
		("No", 6),
		("Asset", 14),
		("Model", 32),
		("Registration Number", 18),
		("Chassis", 22),
		("Engine Number", 18),
		("VIN Number", 20),
		("Location", 18),
		("Driver", 44),
		("Required Licence", 24),
		("Driver Licence Status", 18),
		("Vehicle Licence Status", 18),
		("Addendum Status", 16),
		("Overall Status", 16),
		("Comments", 40),
	]

	header_fill = PatternFill(fill_type="solid", fgColor="FFD9EAF7", bgColor="FFD9EAF7")
	owner_fill = PatternFill(fill_type="solid", fgColor="FF262A76", bgColor="FF262A76")
	owner_font = Font(bold=True, color="FFFFFFFF", size=12)
	header_font = Font(bold=True)

	by_owner = {}

	for row in rows:
		by_owner.setdefault(row.owner or "Unassigned", []).append(row)

	current_row = 1
	last_col_letter = ws.cell(row=1, column=len(columns)).column_letter

	for table_index, owner in enumerate(sorted(by_owner), start=1):
		owner_cell = ws.cell(row=current_row, column=1, value=owner)
		owner_cell.font = owner_font

		for col_idx in range(1, len(columns) + 1):
			ws.cell(row=current_row, column=col_idx).fill = owner_fill

		ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(columns))
		current_row += 1

		header_row = current_row

		for col_idx, (label, width) in enumerate(columns, start=1):
			cell = ws.cell(row=current_row, column=col_idx, value=label)
			cell.font = header_font
			cell.fill = header_fill
			ws.column_dimensions[cell.column_letter].width = width

		current_row += 1

		for i, row in enumerate(by_owner[owner], start=1):
			values = [
				i,
				row.asset,
				row.model or row.asset_name,
				row.registration_number,
				row.chassis,
				row.engine_number,
				row.vin_number,
				row.location or "",
				row.driver_display,
				row.required_licence_type or "",
				row.driver_licence_status,
				row.vehicle_licence_status,
				row.addendum_status,
				row.overall_status,
				row.comments or "",
			]

			for col_idx, value in enumerate(values, start=1):
				ws.cell(row=current_row, column=col_idx, value=value)

			current_row += 1

		# A real Excel Table (not a sheet-wide AutoFilter, which only allows
		# one range per sheet) gives this owner's header row its own
		# independent filter dropdowns, scoped to just its own rows.
		table = Table(
			displayName=_sanitize_table_name(owner, table_index),
			ref=f"A{header_row}:{last_col_letter}{current_row - 1}",
		)
		table.tableStyleInfo = TableStyleInfo(
			name="TableStyleMedium9",
			showRowStripes=True,
			showFirstColumn=False,
			showLastColumn=False,
			showColumnStripes=False,
		)
		ws.add_table(table)

		current_row += 1  # blank spacer row between owners

	out = BytesIO()
	wb.save(out)
	out.seek(0)

	return base64.b64encode(out.read()).decode("utf-8")


class VehicleAllocation(Document):
	def autoname(self):
		""""{asset} - {valid_from}" (the JSON's own "format:" autoname string
		is unused once a controller defines its own autoname() — kept there
		only as documentation) — collision-safe: a second allocation for the
		same Asset on the same date gets "{asset} - {valid_from} - 1", a
		third " - 2", and so on."""
		base_name = f"{self.asset} - {self.valid_from}"

		if not frappe.db.exists("Vehicle Allocation", base_name):
			self.name = base_name
			return

		counter = 1

		while frappe.db.exists("Vehicle Allocation", f"{base_name} - {counter}"):
			counter += 1

		self.name = f"{base_name} - {counter}"

	# ------------------------------------------------------------------
	# All compliance fields below are virtual (is_virtual: 1 in the JSON) —
	# they hold no DB column and are recomputed on every read, straight from
	# the actual source doctypes. That is deliberate: this data is dynamic
	# by nature (a licence can expire, be renewed, or be uploaded at any
	# time regardless of this document), so nothing here is ever cached or
	# allowed to go stale between scheduled refreshes.
	#
	# Driver Licence / Company Vehicle Undertaking / Vehicle Licence
	# Compliance are HTML fields, not individual Select/Date/Link fields —
	# once a vehicle can have more than one Driver, a single status/valid_to
	# pair per section can't represent it (each driver has their own), so
	# they render a small table instead. Only Overall Status stays a plain
	# Select field, since the list view and reports need one filterable
	# value out of it.
	#
	# Note: because these are virtual, they are NOT populated by bulk
	# frappe.get_all()/list-view queries — only by loading the full
	# Document (form view, frappe.get_doc). The report/dashboard/
	# notifications call the same underlying engineering.controllers.
	# fleet_compliance functions directly instead, for exactly this reason.
	# ------------------------------------------------------------------
	def _drivers(self):
		"""Every Employee listed in the Drivers table, deduped. Empty when
		this is a shared resource allocated to a Location with no driver
		assigned at all."""
		return collect_drivers([row.driver for row in (self.drivers or [])])

	@property
	def driver_licence_compliance_html(self):
		return render_driver_licence_html(self._drivers(), self.required_licence_type)

	@property
	def company_vehicle_undertaking_html(self):
		return render_addendum_html(self._drivers())

	@property
	def drivers_overview_html(self):
		return render_drivers_overview_html(self._drivers(), self.required_licence_type)

	@property
	def vehicle_licence_compliance_html(self):
		return render_vehicle_licence_html(self.asset)

	@property
	def overall_status(self):
		return compute_all(self.asset, self._drivers(), self.required_licence_type)["overall_status"]

	@property
	def service_history_html(self):
		return render_service_history_html(self.asset)

	# ------------------------------------------------------------------
	# Allocation-period lifecycle (real, stored fields — these describe
	# this document's own identity/state, not another doctype's data).
	# ------------------------------------------------------------------
	def validate(self):
		self.warn_about_active_conflicts()

	def warn_about_active_conflicts(self):
		"""Non-blocking heads-up, surfaced on every save (Draft or Submit) —
		the form's own JS runs the same checks live on Asset/Drivers change,
		this is the server-side backstop so the warning shows up regardless
		of what the browser did (API/Data Import saves included, though
		nothing displays a msgprint there).

		Both the Asset side and the Driver side are auto-resolved on submit
		now (close_previous_open_allocation /
		close_or_update_previous_driver_allocations) — this is purely an
		FYI about what submitting will do, not a call to action, so both
		read as informational (blue). Mirrors check_asset_conflict/
		check_driver_conflicts in vehicle_allocation.js exactly, so the
		wording is never a surprise between what you saw while editing and
		what shows up on save."""
		conflicts = _find_active_conflicts(self.asset, self._drivers(), exclude_name=self.name)

		if not conflicts:
			return

		def route(allocation):
			return f"/app/vehicle-allocation/{frappe.utils.quote(allocation)}"

		asset_conflicts = [c for c in conflicts if c["type"] == "asset"]
		driver_conflicts = [c for c in conflicts if c["type"] == "employee"]

		if asset_conflicts:
			links = ", ".join(f"<a href='{route(c['allocation'])}'>{c['allocation']}</a>" for c in asset_conflicts)
			frappe.msgprint(
				frappe._(
					"{0} is currently allocated under {1}. Submitting this allocation will automatically close"
					" that one — no action needed."
				).format(frappe.utils.escape_html(self.asset), links),
				title=frappe._("Asset Already Allocated"),
				indicator="blue",
			)

		if driver_conflicts:
			lines = []

			for c in driver_conflicts:
				link = f"<a href='{route(c['allocation'])}'>{c['allocation']}</a>"
				driver_label = f"<b>{frappe.utils.escape_html(c.get('driver_name') or c['driver'])}</b>"

				if c["sole_driver"]:
					lines.append(
						frappe._("{0}'s other allocation ({1}) will be closed automatically — a change of vehicle.").format(
							driver_label, link
						)
					)
				else:
					lines.append(
						frappe._(
							"{0} will be removed from the shared allocation {1} — that vehicle stays allocated to its"
							" other driver(s)."
						).format(driver_label, link)
					)

			frappe.msgprint(
				"<br>".join(lines),
				title=frappe._("Driver Reassignment"),
				indicator="blue",
			)

	def before_submit(self):
		"""Signed Vehicle Handover Paperwork is only meaningful once there is
		someone to hand the vehicle to — required at submit time whenever
		one or more Drivers are listed, never required for a shared vehicle
		with no driver assigned. Deliberately not a plain reqd/
		mandatory_depends_on field: that would also block ordinary Draft
		saves while the allocation is still being assembled, one Driver at
		a time."""
		if self._drivers() and not self.handover_paperwork:
			frappe.throw(
				frappe._(
					"You cannot submit this allocation without attaching the signed Vehicle Handover"
					" Paperwork — one or more Drivers are listed above."
				),
				title=frappe._("Attachment Required"),
			)

	def on_submit(self):
		self.close_previous_open_allocation()
		self.close_or_update_previous_driver_allocations()
		self.sync_asset_custodian_and_location()

	def on_update_after_submit(self):
		"""Fires when an allow_on_submit field changes on an already-
		submitted (Current) allocation — in practice, just Location today.
		Re-run the same reconciliation so moving the vehicle here raises an
		Asset Movement too, not just on the original submit."""
		self.sync_asset_custodian_and_location()

	def close_previous_open_allocation(self):
		"""A newly submitted allocation for an Asset supersedes whichever
		other allocation for that same Asset was still open. valid_to is a
		system-only field — it is never typed in by a user."""
		previous = frappe.get_all(
			"Vehicle Allocation",
			filters={"asset": self.asset, "status": "Current", "docstatus": 1, "name": ["!=", self.name]},
			pluck="name",
		)

		for name in previous:
			frappe.db.set_value(
				"Vehicle Allocation", name, {"status": "Closed", "valid_to": self.valid_from}
			)

	def close_or_update_previous_driver_allocations(self):
		"""Mirrors close_previous_open_allocation, but for Drivers instead
		of the Asset: a driver moving to a new vehicle is a change of
		vehicle for them, same as a new allocation for an Asset supersedes
		the previous one — the default is to close their old allocation
		out, not leave it dangling. A driver CAN legitimately be on two
		vehicles at once though, so this only auto-closes the other
		allocation outright when this driver was its SOLE driver (their
		change of vehicle really did vacate it). If it's a shared vehicle
		(more than one Driver), closing the whole thing would strand
		whoever else is still using it — there, only this driver is
		removed from it; the allocation stays Current for the rest."""
		drivers = self._drivers()

		if not drivers:
			return

		conflicts = [
			c for c in _find_active_conflicts(self.asset, drivers, exclude_name=self.name) if c["type"] == "employee"
		]

		by_parent = {}

		for c in conflicts:
			by_parent.setdefault(c["allocation"], {"sole_driver": c["sole_driver"], "drivers": []})
			by_parent[c["allocation"]]["drivers"].append(c["driver"])

		for parent_name, info in by_parent.items():
			if info["sole_driver"]:
				frappe.db.set_value(
					"Vehicle Allocation", parent_name, {"status": "Closed", "valid_to": self.valid_from}
				)
			else:
				frappe.db.delete(
					"Vehicle Allocation Driver",
					{
						"parent": parent_name,
						"parentfield": "drivers",
						"driver": ["in", info["drivers"]],
					},
				)

	def sync_asset_custodian_and_location(self):
		"""Keep the linked Asset's own Custodian/Location in step with this
		allocation, via a proper Asset Movement (never a silent field
		overwrite — Asset Movement is the audited, standard ERPNext
		mechanism, and other tools already read/rely on it).

		Custodian rule: exactly one Driver -> that Employee is Custodian.
		More than one Driver -> the Location's configured recipient (Fleet
		Management Settings) is Custodian, standing in for "whoever
		manages this shared vehicle at this Location". Zero Drivers -> the
		Asset's existing Custodian is left untouched (a shared-resource
		allocation with nobody named doesn't imply anything about who
		should hold it).

		IMPORTANT: ERPNext's Asset Movement is not a diff — on submit it
		takes the LATEST movement's own target_location/to_employee as the
		Asset's complete new state (see AssetMovement.
		get_latest_location_and_custodian), and unlike location (which is
		only applied if truthy), custodian is applied unconditionally, even
		when blank. So a pure "Transfer" movement that leaves to_employee
		empty silently wipes the Custodian. Every movement created here
		therefore always carries the CURRENT value forward for whichever
		dimension isn't the one actually changing, so each movement is a
		complete, accurate snapshot rather than a partial one.

		Only creates a movement for whichever of Custodian/Location
		actually differs from what the Asset currently records — calling
		this repeatedly (e.g. on every subsequent Location edit) is a
		no-op once things already match."""
		if not self.asset:
			return

		asset = frappe.db.get_value("Asset", self.asset, ["custodian", "location", "company"], as_dict=True)

		if not asset:
			return

		drivers = self._drivers()

		if len(drivers) == 1:
			target_custodian = drivers[0]
		elif len(drivers) == 0:
			target_custodian = asset.custodian  # leave untouched
		else:
			target_custodian = get_location_custodian(self.location) or asset.custodian

		target_location = self.location or asset.location
		current_location = asset.location
		current_custodian = asset.custodian

		if target_location and target_location != current_location:
			self._create_asset_movement(
				purpose="Transfer",
				company=asset.company,
				source_location=current_location,
				target_location=target_location,
				to_employee=current_custodian,  # carry forward — see docstring
			)
			current_location = target_location

		if target_custodian and target_custodian != current_custodian:
			self._create_asset_movement(
				purpose="Issue",
				company=asset.company,
				target_location=current_location,  # carry forward — see docstring
				to_employee=target_custodian,
			)

	def _create_asset_movement(
		self, purpose, company, source_location=None, target_location=None, from_employee=None, to_employee=None
	):
		movement = frappe.new_doc("Asset Movement")
		movement.purpose = purpose
		movement.company = company
		movement.transaction_date = now_datetime()
		movement.reference_doctype = "Vehicle Allocation"
		movement.reference_name = self.name
		movement.append(
			"assets",
			{
				"asset": self.asset,
				"source_location": source_location,
				"target_location": target_location,
				"from_employee": from_employee,
				"to_employee": to_employee,
			},
		)
		movement.insert(ignore_permissions=True)
		movement.submit()

		return movement.name

	def on_cancel(self):
		self.db_set("status", "Cancelled")
		self.cancel_linked_asset_movements()

	def cancel_linked_asset_movements(self):
		"""Any Asset Movement this allocation raised (see
		sync_asset_custodian_and_location) dynamically links back to it via
		reference_doctype/reference_name — Frappe blocks cancelling a
		document that's still linked from elsewhere, so without this a
		Vehicle Allocation that ever changed the Asset's Custodian/Location
		could never be cancelled at all."""
		for name in frappe.get_all(
			"Asset Movement",
			filters={"reference_doctype": "Vehicle Allocation", "reference_name": self.name, "docstatus": 1},
			pluck="name",
		):
			frappe.get_doc("Asset Movement", name).cancel()

	@frappe.whitelist()
	def close(self):
		"""'Return Vehicle': close this allocation (asset goes back to the
		pool) without creating a new allocation record."""
		if self.docstatus != 1:
			frappe.throw(frappe._("Only a submitted allocation can be returned."))

		if self.status != "Current":
			frappe.throw(frappe._("This allocation is already {0}.").format(self.status))

		self.db_set({"status": "Closed", "valid_to": nowdate()})
